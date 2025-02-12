from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func
from ...models.grade import StudentGrade
from ...models.course import Course
from ...models.student import Student
from ...models.student_courses import StudentCourse
from ...core.database import get_db
from typing import Dict, List
from pydantic import BaseModel
from datetime import datetime
import openpyxl
from io import BytesIO
from urllib.parse import quote
from openai import OpenAI
import json
import asyncio

router = APIRouter()

class GradeBase(BaseModel):
    daily_score: float | None = None
    exam_score: float | None = None

class GradeResponse(GradeBase):
    student_id: int
    name: str
    total_score: float | None = None

    class Config:
        from_attributes = True

class GradeStatistics(BaseModel):
    total_students: int
    average_score: float
    max_score: float
    min_score: float
    pass_rate: float
    failing_number: int
    score_distribution: List[Dict]

# 初始化OpenAI客户端
client = OpenAI(
    api_key="sk-Y3BU8GtMz4xUmax0jlHEaSprGooS2rj9KaJoAovnizA0hiRz",
    base_url="https://tbnx.plus7.plus/v1"
)

@router.get("/course/{course_id}/grades")
async def get_course_grades(course_id: int, db: Session = Depends(get_db)) -> Dict:
    try:
        # 获取所有选课学生的成绩
        grades = db.query(
            Student.student_id,
            Student.name,
            StudentGrade.daily_score,
            StudentGrade.exam_score
        ).join(
            StudentGrade, 
            (StudentGrade.student_id == Student.student_id) & 
            (StudentGrade.course_id == course_id),
            isouter=True
        ).join(
            StudentCourse,
            (StudentCourse.student_id == Student.student_id) &
            (StudentCourse.course_id == course_id)
        ).all()
        
        grade_responses = []
        for grade in grades:
            total_score = None
            if grade.daily_score is not None and grade.exam_score is not None:
                total_score = float(grade.daily_score) * 0.4 + float(grade.exam_score) * 0.6
            
            grade_responses.append({
                "student_id": grade.student_id,
                "name": grade.name,
                "daily_score": float(grade.daily_score) if grade.daily_score else None,
                "exam_score": float(grade.exam_score) if grade.exam_score else None,
                "total_score": round(total_score, 2) if total_score is not None else None
            })
        
        return {
            "status": "success",
            "data": grade_responses
        }
    except Exception as e:
        print(f"Error in get_course_grades: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"获取成绩列表失败: {str(e)}"
        )

@router.post("/course/{course_id}/grades")
async def save_grades(
    course_id: int,
    grades: List[GradeBase],
    student_ids: List[int],
    db: Session = Depends(get_db)
) -> Dict:
    try:
        for i, grade_data in enumerate(grades):
            student_id = student_ids[i]
            
            # 检查成绩是否已存在
            grade = db.query(StudentGrade).filter(
                StudentGrade.course_id == course_id,
                StudentGrade.student_id == student_id
            ).first()
            
            if grade:
                # 更新现有成绩
                grade.daily_score = grade_data.daily_score
                grade.exam_score = grade_data.exam_score
            else:
                # 创建新成绩记录
                grade = StudentGrade(
                    course_id=course_id,
                    student_id=student_id,
                    daily_score=grade_data.daily_score,
                    exam_score=grade_data.exam_score
                )
                db.add(grade)
        
        db.commit()
        return {
            "status": "success",
            "message": "成绩保存成功"
        }
    except Exception as e:
        db.rollback()
        print(f"Error in save_grades: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"保存成绩失败: {str(e)}"
        )

@router.get("/course/{course_id}/statistics")
async def get_grade_statistics(course_id: int, db: Session = Depends(get_db)) -> Dict:
    try:
        # 获取成绩统计数据
        grades = db.query(StudentGrade).filter(StudentGrade.course_id == course_id).all()
        
        # 计算总分列表
        total_scores = []
        for grade in grades:
            if grade.daily_score is not None and grade.exam_score is not None:
                # 将 Decimal 转换为 float
                daily = float(grade.daily_score)
                exam = float(grade.exam_score)
                total_score = daily * 0.4 + exam * 0.6
                total_scores.append(total_score)
        
        if not total_scores:
            return {
                "status": "success",
                "data": {
                    "total_students": 0,
                    "average_score": 0,
                    "max_score": 0,
                    "min_score": 0,
                    "pass_rate": "0.0%",
                    "failing_number": 0,
                    "score_distribution": []
                }
            }
        
        # 计算统计数据
        total_students = len(total_scores)
        average_score = sum(total_scores) / total_students
        max_score = max(total_scores)
        min_score = min(total_scores)
        failing_number = sum(1 for score in total_scores if score < 60)
        pass_rate = (total_students - failing_number) / total_students * 100
        
        # 计算分数段分布
        score_ranges = [(0, 59), (60, 69), (70, 79), (80, 89), (90, 100)]
        distribution = []
        
        for start, end in score_ranges:
            count = sum(1 for score in total_scores if start <= score <= end)
            distribution.append({
                "score": f"{start}-{end}",
                "count": count,
                "percentage": f"{(count/total_students*100):.1f}%"
            })
        
        return {
            "status": "success",
            "data": {
                "total_students": total_students,
                "average_score": round(average_score, 1),
                "max_score": round(max_score, 1),
                "min_score": round(min_score, 1),
                "pass_rate": f"{pass_rate:.1f}%",
                "failing_number": failing_number,
                "score_distribution": distribution
            }
        }
    except Exception as e:
        print(f"Error in get_grade_statistics: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"获取成绩统计失败: {str(e)}"
        )

@router.get("/course/{course_id}/export")
async def export_grades(course_id: int, db: Session = Depends(get_db)):
    try:
        # 获取课程信息
        course = db.query(Course).filter(Course.course_id == course_id).first()
        if not course:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="课程不存在"
            )

        # 获取所有学生成绩
        grades = db.query(
            Student.student_id,
            Student.name,
            Student.class_name,
            StudentGrade.daily_score,
            StudentGrade.exam_score
        ).join(
            StudentGrade,
            (StudentGrade.student_id == Student.student_id) &
            (StudentGrade.course_id == course_id),
            isouter=True
        ).join(
            StudentCourse,
            (StudentCourse.student_id == Student.student_id) &
            (StudentCourse.course_id == course_id)
        ).all()

        # 创建工作簿和工作表
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = str(course.name) + "成绩表"

        # 添加课程信息标题行
        ws.merge_cells('A1:F1')
        title_cell = ws.cell(row=1, column=1, value=f"课程：{course.course_code}-{course.name} ({course.semester})")
        title_cell.font = openpyxl.styles.Font(bold=True, size=12)
        title_cell.alignment = openpyxl.styles.Alignment(horizontal='center', vertical='center')
        
        # 设置标题行高
        ws.row_dimensions[1].height = 25

        # 添加表头
        headers = ["学号", "姓名", "班级", "平时成绩(40%)", "考试成绩(60%)", "总成绩"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col, value=header)
            # 设置表头样式
            cell.font = openpyxl.styles.Font(bold=True)
            cell.fill = openpyxl.styles.PatternFill(
                start_color='E0E0E0',
                end_color='E0E0E0',
                fill_type='solid'
            )
            cell.alignment = openpyxl.styles.Alignment(
                horizontal='center',
                vertical='center'
            )
            cell.border = openpyxl.styles.Border(
                left=openpyxl.styles.Side(style='thin'),
                right=openpyxl.styles.Side(style='thin'),
                top=openpyxl.styles.Side(style='thin'),
                bottom=openpyxl.styles.Side(style='thin')
            )

        # 设置表头行高
        ws.row_dimensions[2].height = 20

        # 添加数据
        for row, grade in enumerate(grades, 3):
            cells = [
                (grade.student_id, 'General'),
                (grade.name, 'General'),
                (grade.class_name, 'General'),
                (float(grade.daily_score) if grade.daily_score else None, 'Number'),
                (float(grade.exam_score) if grade.exam_score else None, 'Number'),
                (None, 'Number')  # 总成绩先设为None
            ]
            
            for col, (value, cell_format) in enumerate(cells, 1):
                cell = ws.cell(row=row, column=col, value=value)
                # 设置数据单元格样式
                cell.alignment = openpyxl.styles.Alignment(horizontal='center')
                cell.border = openpyxl.styles.Border(
                    left=openpyxl.styles.Side(style='thin'),
                    right=openpyxl.styles.Side(style='thin'),
                    top=openpyxl.styles.Side(style='thin'),
                    bottom=openpyxl.styles.Side(style='thin')
                )
                if cell_format == 'Number' and value is not None:
                    cell.number_format = '0.0'

            # 计算并设置总成绩
            if grade.daily_score is not None and grade.exam_score is not None:
                total_score = float(grade.daily_score) * 0.4 + float(grade.exam_score) * 0.6
                total_cell = ws.cell(row=row, column=6, value=round(total_score, 1))
                total_cell.number_format = '0.0'
                total_cell.alignment = openpyxl.styles.Alignment(horizontal='center')
                total_cell.border = openpyxl.styles.Border(
                    left=openpyxl.styles.Side(style='thin'),
                    right=openpyxl.styles.Side(style='thin'),
                    top=openpyxl.styles.Side(style='thin'),
                    bottom=openpyxl.styles.Side(style='thin')
                )

        # 设置列宽
        column_widths = {
            'A': 15,  # 学号
            'B': 12,  # 姓名
            'C': 15,  # 班级
            'D': 15,  # 平时成绩
            'E': 15,  # 考试成绩
            'F': 15   # 总成绩
        }
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # 保存到内存中
        excel_file = BytesIO()
        wb.save(excel_file)
        excel_file.seek(0)

        # 生成文件名并进行 URL 编码
        filename = f"{course.name}_{course.semester}_成绩表_{datetime.now().strftime('%Y%m%d')}.xlsx"
        encoded_filename = quote(filename)  # URL 编码文件名

        # 返回文件
        return StreamingResponse(
            excel_file,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                'Content-Disposition': f'attachment; filename="{encoded_filename}"',
                'Access-Control-Expose-Headers': 'Content-Disposition'  # 允许前端访问该头部
            }
        )

    except Exception as e:
        print(f"Error in export_grades: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"导出成绩失败: {str(e)}"
        )

@router.get("/course/{course_id}/ai-analysis")
async def get_ai_analysis(course_id: int, db: Session = Depends(get_db)):
    try:
        # 获取课程信息
        course = db.query(Course).filter(Course.course_id == course_id).first()
        if not course:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="课程不存在"
            )

        # 获取成绩统计数据
        grades = db.query(StudentGrade).filter(StudentGrade.course_id == course_id).all()
        
        # 计算统计数据
        total_scores = []
        for grade in grades:
            if grade.daily_score is not None and grade.exam_score is not None:
                total_score = float(grade.daily_score) * 0.4 + float(grade.exam_score) * 0.6
                total_scores.append(total_score)

        if not total_scores:
            return {
                "status": "success",
                "data": {
                    "analysis": "暂无成绩数据可供分析。"
                }
            }

        # 准备统计数据
        stats = {
            "course_name": course.name,
            "course_code": course.course_code,
            "total_students": len(total_scores),
            "average_score": round(sum(total_scores) / len(total_scores), 1),
            "max_score": round(max(total_scores), 1),
            "min_score": round(min(total_scores), 1),
            "pass_rate": f"{(sum(1 for s in total_scores if s >= 60) / len(total_scores) * 100):.1f}%",
            "score_distribution": {
                "90-100": sum(1 for s in total_scores if 90 <= s <= 100),
                "80-89": sum(1 for s in total_scores if 80 <= s < 90),
                "70-79": sum(1 for s in total_scores if 70 <= s < 80),
                "60-69": sum(1 for s in total_scores if 60 <= s < 70),
                "0-59": sum(1 for s in total_scores if s < 60)
            }
        }

        # 构造AI提示
        prompt = f"""
        作为一个专业的教育分析师，请对以下课程成绩数据进行深入分析，使用Markdown格式：

        课程：{stats['course_name']} ({stats['course_code']})
        总人数：{stats['total_students']}人
        平均分：{stats['average_score']}
        最高分：{stats['max_score']}
        最低分：{stats['min_score']}
        及格率：{stats['pass_rate']}

        分数段分布：
        90-100分：{stats['score_distribution']['90-100']}人
        80-89分：{stats['score_distribution']['80-89']}人
        70-79分：{stats['score_distribution']['70-79']}人
        60-69分：{stats['score_distribution']['60-69']}人
        0-59分：{stats['score_distribution']['0-59']}人

        请从以下几个方面进行分析，使用Markdown格式：

        **总体成绩水平评估：**
        [在这里分析总体成绩水平]

        成绩分布特点分析：
        [在这里分析成绩分布特点]

        存在的问题和不足：
        [在这里分析问题和不足]

        改进建议和教学策略调整方向：
        [在这里提供改进建议]

        请用简洁专业的语言进行分析，每个方面的分析控制在2-3句话。确保"总体成绩水平评估"部分使用加粗格式。
        """

        # 修改 AI 调用为流式响应
        response = client.chat.completions.create(
            model="deepseek-chat",
            messages=[
                {"role": "system", "content": "你是一个专业的教育分析师，擅长分析课程成绩数据并提供建设性建议。"},
                {"role": "user", "content": prompt}
            ],
            temperature=0.7,
            stream=True  # 启用流式响应
        )

        async def generate():
            try:
                for chunk in response:
                    if chunk.choices[0].delta.content is not None:
                        # 使用 Server-Sent Events 格式
                        yield f"data: {json.dumps({'content': chunk.choices[0].delta.content})}\n\n"
                    await asyncio.sleep(0.1)  # 添加小延迟以防止过快
            except Exception as e:
                print(f"Streaming error: {str(e)}")
                yield f"data: {json.dumps({'error': str(e)})}\n\n"

        return StreamingResponse(
            generate(),
            media_type="text/event-stream",
            headers={
                'Cache-Control': 'no-cache',
                'Connection': 'keep-alive',
            }
        )

    except Exception as e:
        print(f"Error in get_ai_analysis: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"AI分析失败: {str(e)}"
        ) 